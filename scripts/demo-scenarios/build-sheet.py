"""Assemble docs/demo-scenarios.xlsx from the scenario run's result log.

Reads .scenario-results/results.jsonl (one row per recorded scenario, append
only) and writes a two-sheet workbook: Scenarios and Summary.

RE-RUNS OVERWRITE, THEY DO NOT ACCUMULATE. The log is append-only so a stage
can be re-run after a fix, and the LAST row recorded for an id wins. A sheet
that carried both the stale row and the fresh one would let a reader quote
whichever supported their conclusion.

If openpyxl is unavailable this writes docs/demo-scenarios.csv instead and
says so on stdout — never silently.
"""

from __future__ import annotations

import json
import os
import sys
from collections import OrderedDict

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
RESULTS = os.path.join(REPO, ".scenario-results", "results.jsonl")
XLSX = os.path.join(REPO, "docs", "demo-scenarios.xlsx")
CSV = os.path.join(REPO, "docs", "demo-scenarios.csv")

COLUMNS = [
    ("id", "ID"),
    ("demoStep", "Demo step"),
    ("scenario", "Scenario"),
    ("surface", "Surface"),
    ("precondition", "Precondition"),
    ("steps", "Steps"),
    ("expected", "Expected"),
    ("actual", "Actual"),
    ("status", "Status"),
    ("evidence", "Evidence"),
    ("notes", "Notes"),
]

# Group order for the sheet: environment first, then the paths in demo order.
GROUP_ORDER = ["S-ENV", "S-BE", "S-CAP", "S-CUI", "S-KDS", "S-ADM", "S-SYNC"]

# Failures ranked by what they cost the demo, most severe first. An id absent
# from this list sorts after every id present in it.
SEVERITY = [
    "S-CUI-03",   # the waiter phone cannot pair at all -> demo step 1a is dead
    "S-SYNC-09",  # the pump's period is why it cannot pair
    "S-ENV-02",   # the POS process vanished mid-run
    "S-SYNC-04",  # gap A7: nothing but `order` reaches the cloud
    "S-ADM-08",   # demo step 4 has no screen
    "S-ADM-09",   # demo step 5 has no screen
    "S-SYNC-03",  # replayed orders carry an unresolvable device_id
    "S-SYNC-08",  # 12 cloud items have no variant, so their lines cannot replay
    "S-KDS-04",
    "S-KDS-05",
    "S-ADM-02",
    "S-BE-09",
    "S-ENV-01",
]


def load_rows():
    if not os.path.exists(RESULTS):
        sys.exit(f"no results at {RESULTS} — run the scenario stages first")
    latest = OrderedDict()
    # utf-8-sig: the log is appended to by Node, but a PowerShell truncation
    # of the file can leave a byte-order mark on the first line.
    with open(RESULTS, "r", encoding="utf-8-sig") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            row = json.loads(line)
            latest[row["id"]] = row  # last write wins
    return list(latest.values())


def sort_key(row):
    rid = row["id"]
    prefix = rid.rsplit("-", 1)[0]
    try:
        group = GROUP_ORDER.index(prefix)
    except ValueError:
        group = len(GROUP_ORDER)
    tail = rid.rsplit("-", 1)[-1]
    return (group, int(tail) if tail.isdigit() else 999, rid)


def severity_key(row):
    try:
        return SEVERITY.index(row["id"])
    except ValueError:
        return len(SEVERITY)


def write_csv(rows):
    import csv

    with open(CSV, "w", encoding="utf-8", newline="") as fh:
        w = csv.writer(fh)
        w.writerow([h for _, h in COLUMNS])
        for r in rows:
            w.writerow([r.get(k, "") for k, _ in COLUMNS])
    print(f"WROTE CSV (openpyxl unavailable): {CSV}")


def write_xlsx(rows, counts, failures):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Scenarios"

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    status_fill = {
        "PASS": PatternFill("solid", fgColor="D1FAE5"),
        "FAIL": PatternFill("solid", fgColor="FEE2E2"),
        "BLOCKED": PatternFill("solid", fgColor="FEF3C7"),
        "NOT TESTABLE": PatternFill("solid", fgColor="E5E7EB"),
    }
    status_font = {
        "PASS": Font(color="065F46", bold=True),
        "FAIL": Font(color="991B1B", bold=True),
        "BLOCKED": Font(color="92400E", bold=True),
        "NOT TESTABLE": Font(color="374151", bold=True),
    }

    ws.append([h for _, h in COLUMNS])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"

    for r in rows:
        ws.append([r.get(k, "") for k, _ in COLUMNS])
        row_idx = ws.max_row
        st = r.get("status", "")
        c = ws.cell(row=row_idx, column=9)
        if st in status_fill:
            c.fill = status_fill[st]
            c.font = status_font[st]
        c.alignment = Alignment(horizontal="center", vertical="top")
        for col in range(1, len(COLUMNS) + 1):
            ws.cell(row=row_idx, column=col).alignment = Alignment(
                vertical="top", wrap_text=col in (3, 5, 6, 7, 8, 10, 11)
            )

    widths = [11, 11, 44, 26, 40, 42, 40, 60, 14, 40, 62]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:K{ws.max_row}"

    # ------------------------------------------------------------- Summary
    s = wb.create_sheet("Summary")
    title = Font(bold=True, size=13)
    s["A1"] = "Holler demo scenario run"
    s["A1"].font = Font(bold=True, size=15)
    s["A2"] = "Every row in Scenarios was executed against the live stack. Nothing is mocked."
    s["A3"] = "A scenario that could not be driven is NOT TESTABLE or BLOCKED, never PASS."

    s["A5"] = "Counts by status"
    s["A5"].font = title
    s.append([])
    r = 6
    for st in ["PASS", "FAIL", "BLOCKED", "NOT TESTABLE"]:
        s.cell(row=r, column=1, value=st).font = Font(bold=True)
        s.cell(row=r, column=2, value=counts.get(st, 0))
        if st in status_fill:
            s.cell(row=r, column=1).fill = status_fill[st]
        r += 1
    s.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    s.cell(row=r, column=2, value=sum(counts.values())).font = Font(bold=True)

    r += 2
    s.cell(row=r, column=1, value="Top failures, most severe first").font = title
    r += 1
    hdr = ["ID", "Demo step", "Scenario", "What actually happened"]
    for i, h in enumerate(hdr, start=1):
        c = s.cell(row=r, column=i, value=h)
        c.fill = header_fill
        c.font = header_font
    r += 1
    for f in failures:
        s.cell(row=r, column=1, value=f["id"])
        s.cell(row=r, column=2, value=f.get("demoStep", ""))
        s.cell(row=r, column=3, value=f["scenario"]).alignment = Alignment(wrap_text=True, vertical="top")
        s.cell(row=r, column=4, value=f["actual"]).alignment = Alignment(wrap_text=True, vertical="top")
        s.cell(row=r, column=1).fill = status_fill["FAIL"]
        r += 1

    for col, w in zip("ABCD", [14, 12, 46, 96]):
        s.column_dimensions[col].width = w

    wb.save(XLSX)
    print(f"WROTE XLSX: {XLSX}")


def main():
    rows = sorted(load_rows(), key=sort_key)
    counts = {}
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    failures = sorted([r for r in rows if r["status"] == "FAIL"], key=severity_key)

    for st in ["PASS", "FAIL", "BLOCKED", "NOT TESTABLE"]:
        print(f"{st:<14} {counts.get(st, 0)}")
    print(f"{'TOTAL':<14} {len(rows)}")

    try:
        write_xlsx(rows, counts, failures)
    except ImportError:
        write_csv(rows)


if __name__ == "__main__":
    main()
